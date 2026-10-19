
import json
import os
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

import openpyxl

from .domain_knowledge import (
    LABEL_DERIVED, LABEL_SOURCE_COL, SUBJECT_KEYS, SUBJECTS, dominant_group,
)

SOURCE_FILE = os.path.join("data", "DATASET.xlsx")
CAREER_SHEET = "DỮ LIỆU NGHỀ NGHIỆP"
STUDENT_SHEET = "DỮ LIỆU ĐIỂM SINH VIÊN"
CAREERS_EXPORT_FILE = os.path.join("data", "careers_from_dataset.json")
STUDENTS_EXPORT_FILE = os.path.join("data", "students.csv")

os.makedirs("data", exist_ok=True)


def _get_worksheet(workbook, sheet_name):
    if sheet_name not in workbook.sheetnames:
        available = ", ".join(workbook.sheetnames)
        raise ValueError(
            f"Không tìm thấy sheet '{sheet_name}' trong {SOURCE_FILE}. "
            f"Các sheet hiện có: {available}"
        )
    return workbook[sheet_name]


def export_careers(wb):
    ws = _get_worksheet(wb, CAREER_SHEET)
    header = [str(c.value).strip() if c.value else c.value for c in ws[1]]
    careers = {}
    for row in ws.iter_rows(min_row=2, max_col=len(header), values_only=True):
        if not row[0]:
            continue
        info = dict(zip(header, row))
        career_id = str(row[0]).strip()
        info["career_group"] = str(info.get("career_group", "")).strip()
        careers[career_id] = info
    with open(CAREERS_EXPORT_FILE, "w", encoding="utf-8") as f:
        json.dump(careers, f, ensure_ascii=False, indent=2)
    print(f"Đã tạo {CAREERS_EXPORT_FILE} ({len(careers)} nghề)")
    return careers


def export_students(wb):
    ws = _get_worksheet(wb, STUDENT_SHEET)
    header = [str(c.value).strip() if c.value else c.value for c in ws[1]]
    # Map tên cột trong Excel (có thể lệch khoảng trắng, vd "Trí tuệ nhân tạo ")
    # về đúng tên chuẩn trong SUBJECTS.
    col_to_subject = {}
    for idx, col_name in enumerate(header):
        if not col_name:
            continue
        stripped = col_name.strip()
        for subj in SUBJECTS:
            if stripped == subj or stripped == subj.strip():
                col_to_subject[idx] = subj
                break

    if "Họ và Tên" in header:
        name_idx = header.index("Họ và Tên")
    elif "STT" in header:
        name_idx = header.index("STT")
    else:
        name_idx = 0

    rows_out = []
    skipped = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        name = row[name_idx] if name_idx < len(row) else ""
        scores = {}
        ok = True
        for idx, subj in col_to_subject.items():
            val = row[idx] if idx < len(row) else None
            if val is None:
                ok = False
                break
            scores[subj] = float(val)
        if not ok or len(scores) != len(SUBJECTS):
            skipped += 1
            continue
        label = dominant_group(scores)
        record = {"name": name}
        for subj in SUBJECTS:
            record[SUBJECT_KEYS[subj]] = scores[subj]
        record["career_group"] = label
        record[LABEL_SOURCE_COL] = LABEL_DERIVED  # nhãn suy ra từ bảng điểm (pseudo-label)
        rows_out.append(record)

    fieldnames = ["name"] + [SUBJECT_KEYS[s] for s in SUBJECTS] + ["career_group", LABEL_SOURCE_COL]
    import csv
    with open(STUDENTS_EXPORT_FILE, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows_out)
    print(f"Đã tạo {STUDENTS_EXPORT_FILE} ({len(rows_out)} sinh viên, bỏ qua {skipped} dòng thiếu điểm)")

    print("\n[LƯU Ý] career_group là NHÃN GIẢ (pseudo-label) suy ra từ chính bảng điểm, không phải")
    print("        nghề nghiệp thực tế của sinh viên. Xem giải thích ở domain_knowledge.py và README.md.")

    # In nhanh phân bố nhãn suy ra được, để kiểm tra dữ liệu có bị lệch nhóm
    # quá mức không (thông tin tham khảo, không ảnh hưởng tới việc huấn luyện).
    from collections import Counter
    dist = Counter(r["career_group"] for r in rows_out)
    print("Phân bố nhãn career_group suy ra từ bảng điểm:")
    for group, count in dist.most_common():
        print(f"  {group:<28}{count:>4}  ({count / len(rows_out) * 100:.1f}%)")
    return rows_out


if __name__ == "__main__":
    wb = openpyxl.load_workbook(SOURCE_FILE, data_only=True)
    export_careers(wb)
    export_students(wb)
    print("Hoàn tất.")
