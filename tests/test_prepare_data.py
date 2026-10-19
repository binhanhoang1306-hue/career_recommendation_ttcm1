import csv
import json
import os
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)

from career_ai import prepare_data
from career_ai.domain_knowledge import SUBJECTS


class TestPrepareData(unittest.TestCase):
    def test_career_export_uses_source_sheet_without_overwriting_runtime_catalog(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = prepare_data.CAREER_SHEET
        sheet.append(("career_id ", "career_name", "career_group"))
        sheet.append(("BE01", "Backend Developer", "Phát triển phần mềm"))

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = os.path.join(temp_dir, "careers_from_dataset.json")
            runtime_catalog_path = os.path.join(temp_dir, "careers.json")
            with open(runtime_catalog_path, "w", encoding="utf-8") as file:
                file.write('{"curated": true}')

            with patch.object(prepare_data, "CAREERS_EXPORT_FILE", output_path):
                careers = prepare_data.export_careers(workbook)

            with open(output_path, encoding="utf-8") as file:
                exported = json.load(file)
            with open(runtime_catalog_path, encoding="utf-8") as file:
                runtime_catalog = file.read()

        self.assertEqual(careers["BE01"]["career_name"], "Backend Developer")
        self.assertEqual(exported, careers)
        self.assertEqual(runtime_catalog, '{"curated": true}')

    def test_missing_sheet_error_lists_available_sheets(self):
        workbook = Workbook()
        workbook.active.title = "Sheet thay thế"

        with self.assertRaisesRegex(ValueError, "Sheet thay thế"):
            prepare_data.export_careers(workbook)

    def test_student_export_uses_serial_column_when_names_are_absent(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = prepare_data.STUDENT_SHEET
        sheet.append(("STT", *SUBJECTS))
        sheet.append((1, *(7.0 for _ in SUBJECTS)))

        with tempfile.TemporaryDirectory() as temp_dir:
            output_path = os.path.join(temp_dir, "students.csv")
            with patch.object(prepare_data, "STUDENTS_EXPORT_FILE", output_path):
                rows = prepare_data.export_students(workbook)

            with open(output_path, newline="", encoding="utf-8-sig") as file:
                exported = list(csv.DictReader(file))

        self.assertEqual(len(rows), 1)
        self.assertEqual(exported[0]["name"], "1")
        self.assertEqual(exported[0]["sec"], "7.0")


if __name__ == "__main__":
    unittest.main()
